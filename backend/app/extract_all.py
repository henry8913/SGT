import json
import sys
import os

STEP_MAP = {
    "Proprietà_beam": "profili",
    "Caratteristiche_macchina": "macchina",
    "Geometria_braccio": "geometria",
    "Masse_proprie": "masse",
    "Baricentri": "baricentri",
    "A_b": "aree_vento",
    "A_rc": "aree_vento",
    "A_cb": "aree_vento",
    "A_Pu": "aree_vento",
    "Vento": "vento",
    "Stabilità C25-Q": "stabilita_q",
    "Stabilità C25-D": "stabilita_d",
    "Carrichi ralla e base - C25": "carichi_ralla",
    "Curve_di_carico II": "curve_carico",
    "Curve_di_carico II IV": "curve_carico",
    "Diagramma di carico": "diagramma",
    "Elenchi a discesa": "dropdown",
}

INCLUDED_SHEETS = set(STEP_MAP.keys())


def col_to_number(col):
    n = 0
    for c in col:
        n = n * 26 + (ord(c) - ord('A') + 1)
    return n


def number_to_col(n):
    s = ""
    while n > 0:
        n, r = divmod(n - 1, 26)
        s = chr(r + ord('A')) + s
    return s


def _is_blue(cell):
    try:
        fill = cell.fill
        if fill and fill.start_color and fill.start_color.rgb:
            rgb = str(fill.start_color.rgb).upper()
            for blue in ['0000FF', '0070C0', '00B0F0', '4472C4', '5B9BD5', 'BDD7EE', '8DB4E2']:
                if blue in rgb:
                    return True
            if fill.start_color.theme is not None:
                if fill.start_color.theme in (4, 5, 6, 7):
                    return True
    except Exception:
        pass
    return False


def extract_all(xlsx_path):
    import openpyxl
    wb = openpyxl.load_workbook(xlsx_path, data_only=False, keep_vba=False)

    all_cells = []
    cell_map = {}

    for sheet_name in wb.sheetnames:
        if sheet_name not in INCLUDED_SHEETS:
            continue
        ws = wb[sheet_name]
        step = STEP_MAP[sheet_name]

        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=min(ws.max_column, 120)):
            for cell in row:
                if cell.value is None:
                    continue
                cell_text = str(cell.value).strip()
                cell_ref = cell.coordinate

                if cell_text.startswith("=") and cell_text != "=":
                    entry = {
                        "step": step,
                        "sheet": sheet_name,
                        "campo": cell_ref,
                        "cell_type": "formula",
                        "formula": cell_text,
                        "default_value": None,
                        "label": "",
                    }
                elif _is_blue(cell):
                    try:
                        val = float(cell.value)
                        entry = {
                            "step": step,
                            "sheet": sheet_name,
                            "campo": cell_ref,
                            "cell_type": "input",
                            "formula": str(val),
                            "default_value": str(val),
                            "label": "",
                        }
                    except (ValueError, TypeError):
                        continue
                elif isinstance(cell.value, (int, float)) and not cell_text.startswith("="):
                    entry = {
                        "step": step,
                        "sheet": sheet_name,
                        "campo": cell_ref,
                        "cell_type": "constant",
                        "formula": str(cell.value),
                        "default_value": str(cell.value),
                        "label": "",
                    }
                else:
                    entry = {
                        "step": step,
                        "sheet": sheet_name,
                        "campo": cell_ref,
                        "cell_type": "label",
                        "formula": cell_text,
                        "default_value": cell_text,
                        "label": cell_text,
                    }

                all_cells.append(entry)
                key = (sheet_name, cell.row, cell.column)
                cell_map[key] = entry

    for entry in all_cells:
        if entry["cell_type"] in ("input", "formula", "constant") and not entry["label"]:
            campo = entry["campo"]
            sheet = entry["sheet"]
            import re
            col_letters = re.match(r'([A-Z]+)', campo).group(1)
            row_num = int(re.search(r'(\d+)', campo).group(1))
            col_num = col_to_number(col_letters)
            best_label = ""
            for c in range(col_num - 1, 0, -1):
                key = (sheet, row_num, c)
                if key in cell_map:
                    neighbor = cell_map[key]
                    if neighbor["cell_type"] == "label":
                        best_label = neighbor["label"]
                        break
            if not best_label:
                for r in range(row_num - 1, max(0, row_num - 10), -1):
                    key = (sheet, r, col_num)
                    if key in cell_map:
                        neighbor = cell_map[key]
                        if neighbor["cell_type"] == "label":
                            best_label = neighbor["label"]
                            break
            entry["label"] = best_label

    wb.close()
    return all_cells


def main():
    if len(sys.argv) < 2:
        print("Usage: python extract_all.py <path_to.xlsm> [output.json]")
        sys.exit(1)
    xlsx_path = sys.argv[1]
    output_path = sys.argv[2] if len(sys.argv) > 2 else "formulas_seed.json"
    print(f"Extracting from: {xlsx_path}")
    cells = extract_all(xlsx_path)
    with open(output_path, "w") as f:
        json.dump(cells, f, indent=2)
    print(f"Extracted {len(cells)} cells -> {output_path}")


if __name__ == "__main__":
    main()
