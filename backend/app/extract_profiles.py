"""Estrae dal foglio Proprietà_beam i profili strutturali e il blocco
"Conversione mm2 - m2 e mm4 - m4" del file Stabilità.xlsm.

Struttura attesa del foglio Proprietà_beam:
  - A1              : titolo "Conversione mm2 - m2 e mm4 - m4"
  - A2/D2/G2/L2     : esempio area (mm2 -> m2)
  - A3/D3/G3/L3     : esempio inerzia (mm4 -> m4)
  - A8/H8/L8/P8     : intestazioni Nomenclatura beam / Riferimento / Proprietà / Valore
  - A{r}            : nome profilo (inizio blocco)
  - H{r}            : Riferimento (numero progressivo)
  - L{r}+P{r}       : coppie Proprietà/Valore (Area trasversale, IY, IZ, HY, BZ)
                      con Valore espresso in m2, m4 e metri
  - AA{r} (+ AD{r}) : peso lineare in kg/m (solo per alcuni tubolari)

I valori vengono convertiti nelle unità di lavoro del DB (mm):
  area x 1e6  -> mm2 ; IY/IZ x 1e12 -> mm4 ; HY/BZ x 1e3 -> mm (dimensione esterna)
"""

import json
import os
import sys

SHEET_NAME = "Proprietà_beam"

PROP_TO_FIELD = {
    "Area trasversale": ("area_mm2", 1e6),
    "IY": ("iy_mm4", 1e12),
    "IZ": ("iz_mm4", 1e12),
    "HY": ("hy_mm", 1e3),
    "BZ": ("bz_mm", 1e3),
}


def _col_value(ws, row, col):
    """Valore di una cella: se è una formula semplice (=+X9+1) ne calcola il risultato."""
    v = ws.cell(row, col).value
    if isinstance(v, str):
        s = v.strip()
        if s.startswith("="):
            import re
            m = re.match(r"^=\+?([A-Z]+)(\d+)\+(\d+)$", s)
            if m:
                other = ws.cell(int(m.group(2)), col).value
                if isinstance(other, (int, float)):
                    return other + int(m.group(3))
        return None
    return v


def extract_profiles(xlsx_path):
    import openpyxl
    wb = openpyxl.load_workbook(xlsx_path, data_only=False, keep_vba=False)
    if SHEET_NAME not in wb.sheetnames:
        wb.close()
        raise ValueError(f"Foglio '{SHEET_NAME}' non trovato nel file")

    ws = wb[SHEET_NAME]

    conversions = []
    title = ws.cell(1, 1).value
    for row, grandezza in ((2, "Area trasversale"), (3, "Momento d'inerzia")):
        ref_mm = ws.cell(row, 1).value
        unita_mm = ws.cell(row, 4).value
        unita_m = ws.cell(row, 12).value
        fattore = 1e6 if "mm2" in str(unita_mm) else 1e12
        if isinstance(ref_mm, (int, float)):
            conversions.append({
                "grandezza": grandezza,
                "unita_mm": str(unita_mm or "mm2"),
                "unita_m": str(unita_m or ("m2" if fattore == 1e6 else "m4")),
                "fattore": fattore,
                "riferimento_mm": float(ref_mm),
                "riferimento_m": float(ref_mm) / fattore,
            })

    profiles = []
    current = None
    last_rif = None
    for row in range(9, ws.max_row + 1):
        nome = ws.cell(row, 1).value
        if nome is not None and str(nome).strip():
            if current:
                profiles.append(current)
            rif = _col_value(ws, row, 8)
            if rif is None and last_rif is not None:
                rif = last_rif + 1
            if isinstance(rif, (int, float)):
                last_rif = int(rif)
            peso = ws.cell(row, 27).value
            current = {
                "nome": str(nome).strip(),
                "riferimento": last_rif,
                "area_mm2": None,
                "iy_mm4": None,
                "iz_mm4": None,
                "hy_mm": None,
                "bz_mm": None,
                "peso_kg_m": float(peso) if isinstance(peso, (int, float)) else None,
            }

        if current is None:
            continue
        prop = ws.cell(row, 12).value
        val = ws.cell(row, 16).value
        if prop is None or not isinstance(val, (int, float)):
            continue
        mapping = PROP_TO_FIELD.get(str(prop).strip())
        if mapping:
            field, factor = mapping
            current[field] = float(val) * factor

    if current:
        profiles.append(current)
    wb.close()

    return {"conversions": conversions, "profiles": profiles}


def main():
    if len(sys.argv) < 2:
        print("Usage: python extract_profiles.py <path_to.xlsm> [output.json]")
        sys.exit(1)
    xlsx_path = sys.argv[1]
    output_path = sys.argv[2] if len(sys.argv) > 2 else os.path.join(
        os.path.dirname(os.path.abspath(__file__)), "beam_profiles_seed.json"
    )
    data = extract_profiles(xlsx_path)
    with open(output_path, "w") as f:
        json.dump(data, f, indent=2, ensure_ascii=False)
    print(f"Extracted {len(data['profiles'])} profiles + {len(data['conversions'])} conversions -> {output_path}")


if __name__ == "__main__":
    main()
